import pandas
from openpyxl import Workbook
from openpyxl.styles import Alignment


class ExcelWriter:
    def __init__(self, path):
        self.path = path
        self.workbook = Workbook()
        self.sheet = self.workbook.active

        path.parent.mkdir(parents=True, exist_ok=True)
        self._init_header()

    @staticmethod
    def _column_names():
        return [
            'Dataset', 'Patient', 'FDI', '3D', '2D',
            *[
                f'{method} {value}°' if isinstance(value, int) else f'{value} {method}'
                for value in [*range(0, 360, 45), 'Mean', 'Max']
                for method in ['Contour', 'Area']
            ],
        ]

    def _init_header(self):
        center_bottom = Alignment(horizontal='center', vertical='bottom')
        center = Alignment(horizontal='center', vertical='center')

        for column, value in enumerate(['Dataset', 'Patient', 'FDI', '3D', '2D'], 1):
            cell = self.sheet.cell(1, column, value)
            cell.alignment = center_bottom
            self.sheet.merge_cells(start_row=1, start_column=column, end_row=2, end_column=column)

        for offset, value in enumerate([*range(0, 360, 45), 'Mean', 'Max']):
            column = 6 + offset * 2
            title = f'{value}°' if isinstance(value, int) else value
            cell = self.sheet.cell(1, column, title)
            cell.alignment = center
            self.sheet.merge_cells(start_row=1, start_column=column, end_row=1, end_column=column + 1)
            for value_column, method in enumerate(['Contour', 'Area'], column):
                cell = self.sheet.cell(2, value_column, method)
                cell.alignment = center

        self.sheet.freeze_panes = 'D3'

    def write(self, dataset, patient, rows):
        center = Alignment(horizontal='center', vertical='center')

        start_row = self.sheet.max_row + 1
        for row, values in enumerate(rows, start_row):
            for column, value in enumerate(values, 3):
                cell = self.sheet.cell(row, column, value)
        end_row = self.sheet.max_row

        for column, value in enumerate([dataset, patient], 1):
            cell = self.sheet.cell(start_row, column, value)
            cell.alignment = center
            self.sheet.merge_cells(start_row=start_row, start_column=column, end_row=end_row, end_column=column)

    def save(self):
        self.workbook.save(self.path)

def load_excel(filename):
    excel = pandas.read_excel(filename, header=None, skiprows=2)
    key_columns = ['Dataset', 'Patient', 'FDI']
    columns = ExcelWriter._column_names()

    if len(excel.columns) - len(columns) == 1:
        columns.append('extra')
    if len(excel.columns) != len(columns):
        raise ValueError(f'expected {len(columns)} columns, got {len(excel.columns)} in "{filename}"')

    excel.columns = columns
    excel[key_columns] = excel[key_columns].ffill()

    return excel
